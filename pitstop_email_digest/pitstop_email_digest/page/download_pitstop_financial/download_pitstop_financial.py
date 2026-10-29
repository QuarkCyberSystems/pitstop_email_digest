import json
from io import BytesIO

import frappe
from erpnext.accounts.doctype.account_group.account_group import get_account_group_doc
from erpnext.accounts.report.financial_statements import get_cost_centers_with_children
from frappe import _
from frappe.utils import add_months, cint, flt, format_date, get_year_start, getdate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHEETS = {
    "balance_sheet": "Summarized Balance Sheet",
    "profit_and_loss": "Summarized Profit and Loss",
    "cash_flow": "Summarized Cash Flow",
    "fixed_assets": "Fixed Assets Statement",
    "changes_in_equity": "Statement of Changes in Equity",
    "mom_trial_balance": "MoM Trial Balance",
    "account_group_trial_balance": "Account Group MoM",
}

# Always included in the workbook, whichever reports are selected
MANDATORY_SHEETS = ["mom_trial_balance", "account_group_trial_balance"]

# Account Group root used by each report sheet: (group_level, report_type)
REPORT_ROOTS = {
    "balance_sheet": ("Report Root", "Balance Sheet"),
    "profit_and_loss": ("Report Root", "Profit and Loss"),
    "cash_flow": ("Report Root", "Cash Flow"),
    "fixed_assets": ("Fixed Asset Root", "Balance Sheet"),
    "changes_in_equity": ("Report Root", "Statement of Changes in Equity"),
}

# Statement columns that are linked to the MoM sheets: fieldname -> MoM period
MOM_LINKS = {
    "balance_sheet": {"actual_display": "closing"},
    "profit_and_loss": {"mtd_actual_display": "month", "ytd_actual_display": "ytd"},
    "cash_flow": {"ytd_actual_display": "ytd"},
}

STATEMENT_HEADER_ROW = 5
MOM_HEADER_ROW = 5
FORMULA_TOLERANCE = 0.05
MAX_FORMULA_TERMS = 250

NUMBER_FORMAT = '#,##0.00;(#,##0.00);"-"'
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


@frappe.whitelist()
def enqueue_download(filters):
    if isinstance(filters, str):
        filters = json.loads(filters)

    filters = frappe._dict(filters)
    validate_filters(filters)

    frappe.enqueue(
        generate_pitstop_financial,
        queue="long",
        timeout=3600 * 3,
        job_name=f"Pitstop Financial {filters.company} {filters.report_date} {frappe.session.user}",
        filters=filters,
        user=frappe.session.user,
    )

    return _(
        "Pitstop Financial is being generated. You will be notified when the file is ready."
    )


def validate_filters(filters):
    if not frappe.has_permission("GL Entry", "report"):
        frappe.throw(
            _("Not permitted to view financial statements"), frappe.PermissionError
        )

    if not filters.get("company"):
        frappe.throw(_("Company is mandatory"))
    if not filters.get("report_date"):
        frappe.throw(_("Report Date is mandatory"))

    filters.sheets = [
        s for s in (filters.get("sheets") or []) if s in REPORT_ROOTS
    ] + MANDATORY_SHEETS


def generate_pitstop_financial(filters, user):
    filters = frappe._dict(filters)

    try:
        content = build_workbook(filters)

        file_doc = frappe.get_doc(
            {
                "doctype": "File",
                "file_name": "Pitstop Financial - {0} - {1}.xlsx".format(
                    filters.company, format_date(filters.report_date, "MMM - yyyy")
                ),
                "is_private": 1,
                "content": content,
            }
        )
        # The job runs as the requesting user, so the private file is owned by and downloadable for them
        file_doc.flags.ignore_permissions = True
        file_doc.insert()
        frappe.db.commit()

        notify_user(
            user,
            _("Pitstop Financial for {0} is ready to download").format(
                format_date(filters.report_date)
            ),
            file_doc=file_doc,
        )
    except Exception:
        frappe.db.rollback()
        frappe.log_error(title="Pitstop Financial Download Failed")
        notify_user(
            user, _("Pitstop Financial generation failed. Please check the Error Log.")
        )


def notify_user(user, subject, file_doc=None):
    notification = frappe.get_doc(
        {
            "doctype": "Notification Log",
            "for_user": user,
            "from_user": user,
            "type": "Alert",
            "subject": subject,
            "document_type": "File" if file_doc else None,
            "document_name": file_doc.name if file_doc else None,
            "link": file_doc.file_url if file_doc else None,
        }
    )
    notification.insert(ignore_permissions=True)
    frappe.db.commit()

    frappe.publish_realtime(
        "pitstop_financial_ready",
        {
            "message": subject,
            "file_url": file_doc.file_url if file_doc else None,
            "file_name": file_doc.file_name if file_doc else None,
        },
        user=user,
    )


def build_workbook(filters):
    report_filters = frappe._dict(
        {
            "company": filters.company,
            "report_date": filters.report_date,
            "cost_center": filters.get("cost_center"),
            "tree_view": cint(filters.get("tree_view")),
        }
    )

    wb = Workbook()
    wb.remove(wb.active)

    # The MoM sheets are laid out first so that the statements can reference their cells
    mom = MoMTrialBalance(filters)

    if "balance_sheet" in filters.sheets:
        from erpnext.accounts.report.summarized_balance_sheet.summarized_balance_sheet import (
            execute,
        )

        columns, data = execute(report_filters.copy())
        write_statement_sheet(
            wb,
            SHEETS["balance_sheet"],
            _("BALANCE SHEET AS AT {0}"),
            columns,
            data,
            filters,
            mom=mom,
            mom_links=MOM_LINKS["balance_sheet"],
        )

    if "profit_and_loss" in filters.sheets:
        from erpnext.accounts.report.summarized_profit_and_loss.summarized_profit_and_loss import (
            execute,
        )

        columns, data = execute(report_filters.copy())
        write_statement_sheet(
            wb,
            SHEETS["profit_and_loss"],
            _("PROFIT & LOSS ACCOUNT FOR THE PERIOD ENDED {0}"),
            columns,
            data,
            filters,
            mom=mom,
            mom_links=MOM_LINKS["profit_and_loss"],
        )

    if "cash_flow" in filters.sheets:
        from erpnext.accounts.report.summarized_cash_flow.summarized_cash_flow import (
            execute,
        )

        columns, data = execute(report_filters.copy())
        write_statement_sheet(
            wb,
            SHEETS["cash_flow"],
            _("CASH FLOW STATEMENT FOR THE PERIOD ENDED {0}"),
            columns,
            data,
            filters,
            mom=mom,
            mom_links=MOM_LINKS["cash_flow"],
        )

    if "fixed_assets" in filters.sheets:
        from erpnext.accounts.report.fixed_assets_statement.fixed_assets_statement import (
            execute,
        )

        columns, data = execute(report_filters.copy())
        write_statement_sheet(
            wb,
            SHEETS["fixed_assets"],
            _("FIXED ASSETS STATEMENT AS AT {0}"),
            columns,
            data,
            filters,
            label_field="description",
            mom=mom,
            is_fixed_asset=True,
        )

    if "changes_in_equity" in filters.sheets:
        from erpnext.accounts.report.statement_of_changes_in_equity.statement_of_changes_in_equity import (
            execute,
        )

        # The report is always for the whole company, it has no Cost Center filter
        company_filters = frappe._dict({**filters, "cost_center": None})
        columns, data = execute(frappe._dict({**report_filters, "cost_center": None}))
        write_statement_sheet(
            wb,
            SHEETS["changes_in_equity"],
            _("STATEMENT OF CHANGES IN EQUITY FOR THE PERIOD ENDED {0}"),
            columns,
            data,
            company_filters,
            mom=mom,
            total_field="total_display",
        )

    mom.write_account_sheet(wb)
    mom.write_group_sheet(wb)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def write_title(ws, title, filters, width):
    ws.cell(row=1, column=1, value=filters.company).font = Font(bold=True, size=14)
    ws.cell(
        row=2, column=1, value=title.format(format_date(filters.report_date))
    ).font = Font(bold=True, size=12)
    if filters.get("cost_center"):
        ws.cell(
            row=3, column=1, value=_("Cost Center: {0}").format(filters.cost_center)
        ).font = Font(italic=True)

    for row in (1, 2):
        ws.merge_cells(
            start_row=row, start_column=1, end_row=row, end_column=max(width, 1)
        )


def write_header(ws, row_idx, labels):
    for col_idx, label in enumerate(labels, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[row_idx].height = 32


def write_statement_sheet(
    wb,
    sheet_name,
    title,
    columns,
    data,
    filters,
    label_field="account_display",
    mom=None,
    mom_links=None,
    is_fixed_asset=False,
    total_field=None,
):
    ws = wb.create_sheet(sheet_name)
    value_columns = [c for c in columns if c.get("fieldname") != label_field]
    data = [frappe._dict(d) for d in data]

    write_title(ws, title, filters, len(columns))
    header_row = STATEMENT_HEADER_ROW
    write_header(
        ws, header_row, [_("Particulars")] + [c.get("label") for c in value_columns]
    )

    if is_fixed_asset:
        formulas = get_fixed_asset_formulas(data, value_columns, header_row, mom)
    else:
        formulas = get_statement_formulas(
            data, value_columns, header_row, mom, mom_links, total_field
        )

    row_idx = header_row
    for d in data:
        row_idx += 1

        is_bold = cint(d.is_bold)
        label_cell = ws.cell(row=row_idx, column=1, value=d.get(label_field) or "")
        label_cell.alignment = Alignment(indent=cint(d.indent) * 2)
        if is_bold:
            label_cell.font = Font(bold=True)

        value_type = d.value_type or "Currency"
        for col_idx, column in enumerate(value_columns, 2):
            value = d.get(column["fieldname"])
            if value is None or d.row_type == "Section Break":
                continue

            cell = ws.cell(row=row_idx, column=col_idx)
            if value_type == "Data":
                cell.value = value
            else:
                cell.value = formulas.get((row_idx, column["fieldname"])) or flt(
                    value, 2
                )
                cell.number_format = get_number_format(value_type, d.format_precision)

            if d.right_align or value_type == "Data":
                cell.alignment = Alignment(horizontal="right")
            if is_bold:
                cell.font = Font(bold=True)

        if d.row_type == "Section Total":
            for col_idx in range(1, len(value_columns) + 2):
                ws.cell(row=row_idx, column=col_idx).border = Border(
                    top=THIN, bottom=THIN
                )

    ws.column_dimensions["A"].width = 55 if label_field == "account_display" else 35
    for col_idx in range(2, len(value_columns) + 2):
        ws.column_dimensions[get_column_letter(col_idx)].width = 20

    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    ws.sheet_view.showGridLines = False


def get_number_format(value_type, precision=None):
    precision = (
        cint(precision)
        if precision not in (None, "")
        else (2 if value_type in ("Currency", "Float", "Percent") else 0)
    )
    decimals = ("." + "0" * precision) if precision else ""

    if value_type == "Percent":
        return f'0{decimals}"%"'
    if value_type == "Int":
        return '#,##0;(#,##0);"-"'

    return f'#,##0{decimals};(#,##0{decimals});"-"'


def is_equal(a, b):
    return abs(flt(a) - flt(b)) < FORMULA_TOLERANCE


def signed_sum_formula(terms):
    """terms: [(sign, cell reference)] -> '=A1-B2+...'"""
    if not terms or len(terms) > MAX_FORMULA_TERMS:
        return None

    formula = "".join(f"{'-' if sign < 0 else '+'}{ref}" for sign, ref in terms)
    return "=" + formula.lstrip("+")


def get_statement_formulas(
    data, value_columns, header_row, mom, mom_links, total_field=None
):
    """Excel formulas for a summarized statement, keyed by (row number, fieldname).

    A formula is used only when it reproduces the report's value, otherwise the cell keeps the value:
    - Account and Account Group rows link to the MoM Trial Balance / Account Group MoM sheets
    - Section Total rows add up the rows of their section (or all rows above for a running total)
    - the total column, if any, adds up the columns before it
    """
    formulas = {}
    fieldnames = [c["fieldname"] for c in value_columns]
    letters = {f: get_column_letter(i) for i, f in enumerate(fieldnames, 2)}
    row_numbers = [header_row + 1 + i for i in range(len(data))]

    for d, row_no in zip(data, row_numbers):
        for fieldname, period in (mom_links or {}).items():
            if fieldname not in letters:
                continue

            if d.row_type == "Account" and d.account:
                ref, value = mom.get_account_ref(d.account, period)
            elif d.row_type == "Account Group" and d.account_group:
                ref, value = mom.get_group_ref(d.account_group, period)
            else:
                continue

            if ref is None:
                continue

            for sign in (1, -1):
                if is_equal(sign * value, d.get(fieldname)):
                    formulas[(row_no, fieldname)] = signed_sum_formula([(sign, ref)])
                    break

    multipliers = [get_display_multiplier(d, fieldnames) for d in data]
    section_start = 0
    for i, d in enumerate(data):
        if d.row_type != "Section Total":
            continue

        # Section Total rows either total their own section or everything above (Running Total)
        for start in (section_start, 0):
            included = [j for j in range(start, i) if is_totalled_row(data[j])]
            if not included:
                continue

            terms = [(multipliers[i] * multipliers[j], j) for j in included]
            if all(
                d.get(f) is None
                or is_equal(d.get(f), sum(s * flt(data[j].get(f)) for s, j in terms))
                for f in fieldnames
            ):
                for f in fieldnames:
                    if d.get(f) is not None:
                        formula = signed_sum_formula(
                            [(s, f"{letters[f]}{row_numbers[j]}") for s, j in terms]
                        )
                        if formula:
                            formulas[(row_numbers[i], f)] = formula
                break

        section_start = i + 1

    # A total column (Statement of Changes in Equity) adds up the columns before it
    if total_field in letters and fieldnames.index(total_field) > 0:
        total_idx = fieldnames.index(total_field)
        first, last = letters[fieldnames[0]], letters[fieldnames[total_idx - 1]]
        for d, row_no in zip(data, row_numbers):
            if d.get(total_field) is not None and is_equal(
                d.get(total_field), sum(flt(d.get(f)) for f in fieldnames[:total_idx])
            ):
                formulas[
                    (row_no, total_field)
                ] = f"=SUM({first}{row_no}:{last}{row_no})"

    return formulas


def is_totalled_row(d):
    if d.is_child:
        return False
    if d.row_type in ("Account", "Account Group"):
        return True
    return d.row_type == "Formula" and (d.value_type or "Currency") == "Currency"


def get_display_multiplier(d, fieldnames):
    """The sign the report applied to show this row (display = raw * multiplier)."""
    for f in fieldnames:
        raw = flt(d.get(f[: -len("_display")])) if f.endswith("_display") else 0
        if raw and d.get(f) is not None:
            return 1 if flt(d.get(f)) * raw > 0 else -1

    return 1


def get_fixed_asset_formulas(data, value_columns, header_row, mom):
    """Fixed Assets Statement: openings link to the MoM sheets, closings and NBV are calculated in the sheet."""
    formulas = {}
    row_map = {}
    for i, d in enumerate(data):
        if d.get("fieldname"):
            row_map[d.fieldname] = (header_row + 1 + i, d)

    closing_rules = {
        "asset_closing": [
            (1, "asset_opening"),
            (1, "asset_additions"),
            (1, "asset_disposals"),
            (1, "asset_transfers"),
        ],
        "dep_closing": [
            (1, "dep_opening"),
            (1, "dep_charges"),
            (1, "dep_disposals"),
            (1, "dep_transfers"),
        ],
        "nbv_closing": [(1, "asset_closing"), (-1, "dep_closing")],
        "nbv_opening": [(1, "asset_opening"), (-1, "dep_opening")],
    }

    key_columns = []
    for col_idx, column in enumerate(value_columns, 2):
        fieldname = column["fieldname"]
        letter = get_column_letter(col_idx)

        if fieldname == "total":
            # Total is the sum of every column before it
            for row_no, d in row_map.values():
                if key_columns and is_equal(
                    d.get("total"), sum(flt(d.get(k)) for k, _l in key_columns)
                ):
                    formulas[
                        (row_no, "total")
                    ] = f"=SUM({key_columns[0][1]}{row_no}:{key_columns[-1][1]}{row_no})"
            continue

        key_columns.append((fieldname, letter))

        ref = value = None
        if column.get("account"):
            ref, value = mom.get_account_ref(column["account"], "opening")
        elif column.get("account_group"):
            ref, value = mom.get_group_ref(column["account_group"], "opening")

        if ref:
            for opening_field in ("asset_opening", "dep_opening"):
                if opening_field not in row_map:
                    continue
                row_no, d = row_map[opening_field]
                for sign in (1, -1):
                    if flt(value) and is_equal(sign * value, d.get(fieldname)):
                        formulas[(row_no, fieldname)] = signed_sum_formula(
                            [(sign, ref)]
                        )
                        break

        for target, parts in closing_rules.items():
            if target not in row_map or not all(p in row_map for _s, p in parts):
                continue
            row_no, d = row_map[target]
            if is_equal(
                d.get(fieldname),
                sum(s * flt(row_map[p][1].get(fieldname)) for s, p in parts),
            ):
                formulas[(row_no, fieldname)] = signed_sum_formula(
                    [(s, f"{letter}{row_map[p][0]}") for s, p in parts]
                )

    return formulas


class MoMTrialBalance:
    """Month on month movement of every ledger linked to the selected statements' Account Groups.

    Values are Debit - Credit. Profit and Loss accounts have no opening balance, and Period Closing Vouchers of the
    current year are ignored so that December movements are not reversed, the same way the summarized reports do.

    Value columns are: Year Opening, the 12 months of the year, Closing (YTD).
    """

    account_sheet = SHEETS["mom_trial_balance"]
    group_sheet = SHEETS["account_group_trial_balance"]
    account_label_count = 5
    group_label_count = 2

    def __init__(self, filters):
        self.filters = filters
        self.report_date = getdate(filters.report_date)
        self.year_start = get_year_start(self.report_date)

        self.months = [add_months(self.year_start, i) for i in range(12)]
        self.report_month_idx = (
            (self.report_date.year - self.year_start.year) * 12
            + self.report_date.month
            - self.year_start.month
            + 1
        )

        self.group_tree = []  # (group_name, indent, root) in display order
        self.group_accounts = {}  # group_name -> set of accounts under it at any depth
        self.account_group_options = {}  # account -> [(root, statement group, account group)]

        self.load_accounts()
        self.load_account_groups()
        self.load_balances()
        self.setup_layout()

    def get_roots(self):
        # The primary statements always classify the ledgers, plus the trees of any other selected report
        report_sheets = ["balance_sheet", "profit_and_loss"] + [
            s
            for s in REPORT_ROOTS
            if s in self.filters.sheets
            and s not in ("balance_sheet", "profit_and_loss")
        ]

        roots = []
        for sheet in report_sheets:
            group_level, report_type = REPORT_ROOTS[sheet]
            root = frappe.db.get_value(
                "Account Group",
                {
                    "company": self.filters.company,
                    "group_level": group_level,
                    "report_type": report_type,
                },
                "name",
            )
            if root and root not in roots:
                roots.append(root)

        return roots

    def load_account_groups(self):
        self.displayed_groups = set()
        for root in self.get_roots():
            self.walk_group(root, [], 0, display=True)

        self.account_groups = {}  # account -> (statement group, account group)
        for account, options in self.account_group_options.items():
            self.account_groups[account] = self.get_preferred_group(account, options)

    def walk_group(self, group_name, chain, indent, display):
        if group_name in chain:
            return

        chain = chain + [group_name]

        # A group nested under more than one statement (e.g. the P&L tree inside the Balance Sheet) is listed once
        display = display and group_name not in self.displayed_groups
        if display:
            self.displayed_groups.add(group_name)
            self.group_tree.append((group_name, indent, chain[0]))

        for g in chain:
            self.group_accounts.setdefault(g, set())

        group = get_account_group_doc(group_name, cache="local")
        for row in group.rows:
            if row.row_type == "Account" and row.account:
                for g in chain:
                    self.group_accounts[g].add(row.account)

                statement_group = chain[1] if len(chain) > 1 else chain[0]
                self.account_group_options.setdefault(row.account, []).append(
                    (chain[0], statement_group, chain[-1])
                )

            elif row.row_type == "Account Group" and row.account_group:
                self.walk_group(row.account_group, chain, indent + 1, display)

    def get_preferred_group(self, account, options):
        """Classify an account by the statement root of its own report type (P&L accounts under the P&L root)."""
        report_type = self.account_map.get(account, {}).get("report_type")
        for root, statement_group, account_group in options:
            root_report_type, group_level = frappe.get_cached_value(
                "Account Group", root, ["report_type", "group_level"]
            )
            if group_level == "Report Root" and root_report_type == report_type:
                return statement_group, account_group

        return options[0][1], options[0][2]

    def load_accounts(self):
        self.accounts = frappe.get_all(
            "Account",
            filters={"company": self.filters.company, "is_group": 0},
            fields=[
                "name",
                "account_number",
                "account_name",
                "root_type",
                "report_type",
            ],
            order_by="lft",
        )
        self.account_map = {d.name: d for d in self.accounts}

    def load_balances(self):
        # The cost center condition is always in the query and is switched off by filter_cost_center = 0
        cost_centers = (
            get_cost_centers_with_children(self.filters.cost_center)
            if self.filters.get("cost_center")
            else []
        )
        args = {
            "company": self.filters.company,
            "year_start": self.year_start,
            "report_date": self.report_date,
            "filter_cost_center": 1 if cost_centers else 0,
            "cost_center": cost_centers or [""],
        }

        self.opening = {}
        for d in frappe.db.sql(
            """
			select gle.account, sum(gle.debit) - sum(gle.credit) as balance
			from `tabGL Entry` gle
			inner join `tabAccount` acc on acc.name = gle.account
			where gle.company = %(company)s
				and gle.posting_date < %(year_start)s
				and acc.report_type = 'Balance Sheet'
				and (%(filter_cost_center)s = 0 or gle.cost_center in %(cost_center)s)
			group by gle.account
		""",
            args,
            as_dict=1,
        ):
            self.opening[d.account] = flt(d.balance, 2)

        self.movement = {}
        for d in frappe.db.sql(
            """
			select account, year(posting_date) as year, month(posting_date) as month,
				sum(debit) - sum(credit) as balance
			from `tabGL Entry`
			where company = %(company)s
				and posting_date between %(year_start)s and %(report_date)s
				and voucher_type != 'Period Closing Voucher'
				and (%(filter_cost_center)s = 0 or cost_center in %(cost_center)s)
			group by account, year(posting_date), month(posting_date)
		""",
            args,
            as_dict=1,
        ):
            self.movement.setdefault(d.account, {})[(d.year, d.month)] = flt(
                d.balance, 2
            )

    def get_values(self, accounts):
        opening = sum(self.opening.get(a, 0) for a in accounts)
        months = []
        for m in self.months:
            months.append(
                sum(
                    self.movement.get(a, {}).get((m.year, m.month), 0) for a in accounts
                )
            )

        return [flt(v, 2) for v in [opening] + months + [opening + sum(months)]]

    def get_value_labels(self):
        return (
            [_("Year Opening Balance")]
            + [format_date(m, "MMM-yy") for m in self.months]
            + [_("Closing Balance (YTD)")]
        )

    def get_trial_balance_accounts(self):
        """Linked accounts, plus any unlinked account with a balance so the Trial Balance still sums to zero."""
        out = []
        for d in self.accounts:
            if d.name in self.account_groups:
                out.append((d, *self.account_groups[d.name]))
            elif d.name in self.opening or d.name in self.movement:
                values = self.get_values([d.name])
                if any(values):
                    out.append((d, _("Not Linked"), _("Not Linked")))

        return out

    def setup_layout(self):
        self.trial_balance_accounts = self.get_trial_balance_accounts()
        self.account_rows = {
            d.name: MOM_HEADER_ROW + 1 + i
            for i, (d, _g, _ag) in enumerate(self.trial_balance_accounts)
        }
        self.group_rows = {
            g: MOM_HEADER_ROW + 1 + i for i, (g, _i, _r) in enumerate(self.group_tree)
        }

        self.account_values = {a: self.get_values([a]) for a in self.account_rows}
        self.group_values = {
            g: self.get_values(self.group_accounts.get(g) or [])
            for g in self.group_rows
        }

    # References from other sheets: period is opening / month (report month) / ytd (year to report month) / closing
    def get_account_ref(self, account, period):
        if account not in self.account_rows:
            return None, 0
        return (
            self.get_period_ref(
                self.account_sheet,
                self.account_label_count,
                self.account_rows[account],
                period,
            ),
            self.get_period_value(self.account_values[account], period),
        )

    def get_group_ref(self, group, period):
        if group not in self.group_rows:
            return None, 0
        return (
            self.get_period_ref(
                self.group_sheet, self.group_label_count, self.group_rows[group], period
            ),
            self.get_period_value(self.group_values[group], period),
        )

    def get_period_ref(self, sheet, label_count, row_no, period):
        def col(idx):
            return get_column_letter(label_count + 1 + idx)

        sheet = f"'{sheet}'"
        if period == "opening":
            return f"{sheet}!{col(0)}{row_no}"
        elif period == "month":
            return f"{sheet}!{col(self.report_month_idx)}{row_no}"
        elif period == "ytd":
            return f"SUM({sheet}!{col(1)}{row_no}:{col(self.report_month_idx)}{row_no})"
        else:
            return f"{sheet}!{col(13)}{row_no}"

    def get_period_value(self, values, period):
        if period == "opening":
            return values[0]
        elif period == "month":
            return values[self.report_month_idx]
        elif period == "ytd":
            return sum(values[1 : self.report_month_idx + 1])
        else:
            return values[13]

    def write_value_cells(self, ws, row_no, label_count, values, formulas=None):
        """Opening and months as values (or formulas), Closing = Opening + SUM(months)."""
        first = get_column_letter(label_count + 1)
        for idx, value in enumerate(values[:13]):
            cell = ws.cell(
                row=row_no,
                column=label_count + 1 + idx,
                value=(formulas or {}).get(idx) or value,
            )
            cell.number_format = NUMBER_FORMAT

        month_start, month_end = (
            get_column_letter(label_count + 2),
            get_column_letter(label_count + 13),
        )
        cell = ws.cell(
            row=row_no,
            column=label_count + 14,
            value=f"={first}{row_no}+SUM({month_start}{row_no}:{month_end}{row_no})",
        )
        cell.number_format = NUMBER_FORMAT

    def write_account_sheet(self, wb):
        ws = wb.create_sheet(self.account_sheet)
        value_labels = self.get_value_labels()
        labels = [
            _("Account Number"),
            _("Account Name"),
            _("Root Type"),
            _("Group"),
            _("Account Group"),
        ] + value_labels
        label_count = self.account_label_count

        write_title(
            ws, _("MONTH ON MONTH TRIAL BALANCE AS AT {0}"), self.filters, len(labels)
        )
        header_row = MOM_HEADER_ROW
        write_header(ws, header_row, labels)

        row_idx = header_row
        for account, statement_group, account_group in self.trial_balance_accounts:
            row_idx += 1

            ws.cell(row=row_idx, column=1, value=account.account_number or "")
            ws.cell(row=row_idx, column=2, value=account.account_name)
            ws.cell(row=row_idx, column=3, value=account.root_type)
            ws.cell(row=row_idx, column=4, value=statement_group)
            ws.cell(row=row_idx, column=5, value=account_group)
            self.write_value_cells(
                ws, row_idx, label_count, self.account_values[account.name]
            )

        self.write_total_row(
            ws, row_idx + 1, label_count, len(value_labels), header_row
        )

        for col, width in zip("ABCDE", (16, 45, 12, 35, 40)):
            ws.column_dimensions[col].width = width
        for col_idx in range(label_count + 1, len(labels) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 16

        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(labels))}{row_idx}"
        ws.freeze_panes = ws.cell(row=header_row + 1, column=label_count + 1)

    def write_group_sheet(self, wb):
        ws = wb.create_sheet(self.group_sheet)
        value_labels = self.get_value_labels()
        labels = [_("Account Group"), _("Statement")] + value_labels
        label_count = self.group_label_count

        write_title(
            ws, _("ACCOUNT GROUP MONTH ON MONTH AS AT {0}"), self.filters, len(labels)
        )
        header_row = MOM_HEADER_ROW
        write_header(ws, header_row, labels)

        row_idx = header_row
        for group_name, indent, root in self.group_tree:
            row_idx += 1
            is_root = indent == 0

            label_cell = ws.cell(row=row_idx, column=1, value=group_name)
            label_cell.alignment = Alignment(indent=indent * 2)
            ws.cell(row=row_idx, column=2, value=root)
            self.write_value_cells(
                ws,
                row_idx,
                label_count,
                self.group_values[group_name],
                self.get_group_formulas(group_name),
            )

            if is_root:
                for col_idx in range(1, len(labels) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = Font(bold=True)
                    cell.fill = TOTAL_FILL

        ws.column_dimensions["A"].width = 55
        ws.column_dimensions["B"].width = 25
        for col_idx in range(label_count + 1, len(labels) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 16

        ws.freeze_panes = ws.cell(row=header_row + 1, column=label_count + 1)

    def get_group_formulas(self, group_name):
        """Group = its own accounts (MoM Trial Balance) + its child groups (this sheet), used only when that adds up
        to the group's value, i.e. no account is counted twice through two child groups."""
        group = get_account_group_doc(group_name, cache="local")

        parts = []  # (sheet, label_count, row, values)
        for row in group.rows:
            if row.row_type == "Account" and row.account in self.account_rows:
                parts.append(
                    (
                        self.account_sheet,
                        self.account_label_count,
                        self.account_rows[row.account],
                        self.account_values[row.account],
                    )
                )
            elif (
                row.row_type == "Account Group"
                and row.account_group in self.group_rows
                and row.account_group != group_name
            ):
                parts.append(
                    (
                        None,
                        self.group_label_count,
                        self.group_rows[row.account_group],
                        self.group_values[row.account_group],
                    )
                )
            elif row.row_type in ("Account", "Account Group"):
                return {}

        values = self.group_values[group_name]
        formulas = {}
        for idx in range(13):
            if not is_equal(values[idx], sum(p[3][idx] for p in parts)):
                return {}

            formula = signed_sum_formula(
                [
                    (
                        1,
                        (f"'{sheet}'!" if sheet else "")
                        + f"{get_column_letter(label_count + 1 + idx)}{row_no}",
                    )
                    for sheet, label_count, row_no, _v in parts
                ]
            )
            if not formula:
                return {}
            formulas[idx] = formula

        return formulas

    @staticmethod
    def write_total_row(ws, row_idx, label_count, value_count, header_row):
        ws.cell(row=row_idx, column=1, value=_("Total"))
        for col_idx in range(label_count + 1, label_count + value_count + 1):
            col = get_column_letter(col_idx)
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=f"=SUBTOTAL(9,{col}{header_row + 1}:{col}{row_idx - 1})",
            )
            ws.cell(row=row_idx, column=col_idx).number_format = NUMBER_FORMAT

        for col_idx in range(1, label_count + value_count + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(bold=True)
            cell.fill = TOTAL_FILL
            cell.border = BORDER
